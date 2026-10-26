#!/usr/bin/env python3
"""Build the selected C / Monograph direction as editable PPTX and vector PDF.

Run with PYTHONPATH=/tmp/leadprocure-presentation-libs python3 build_presentation.py.
Uses audited local copy, licensed bundled fonts, and one generated cover artwork.
No application build, production access or screenshots.
"""
from pathlib import Path
import json, math, hashlib
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
import pymupdf

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'output'
PREVIEW = OUT / 'preview'
OUT.mkdir(exist_ok=True)
PREVIEW.mkdir(exist_ok=True)
W, H = 1280, 720
BLACK, WHITE = '161616', 'FFFFFF'
SILVER, MID, MUTED, LIGHT, HAIR = 'B5B5B5', '737373', '525252', 'D1D1D1', 'D6D6D6'
FONT_FILES = {'Display': 'BarlowCondensed-Bold.ttf', 'DisplayRegular': 'BarlowCondensed-Regular.ttf',
              'Body': 'IBMPlexSans-Regular.ttf', 'BodyBold': 'IBMPlexSans-Bold.ttf'}
for name, filename in FONT_FILES.items():
    pdfmetrics.registerFont(TTFont(name, str(ROOT/'fonts'/filename)))

CONTENT = json.loads((ROOT/'content-reviewed.json').read_text())['slides']

def az_upper(s):
    return s.replace('i', 'İ').replace('ı', 'I').upper()

def rgb(s): return RGBColor.from_string(s)

class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(W), Pt(H)
        self.prs.core_properties.title = 'Leadprocure — Korporativ təqdimat / Monograph'
        self.prs.core_properties.subject = 'Azərbaycan dilində korporativ müştəri təqdimatı'
        self.prs.core_properties.author = 'Leadprocure'
        self.pdf = canvas.Canvas(str(OUT/'Leadprocure_Korporativ_Teqdimat_AZ_v2.pdf'), pagesize=(W,H))
        self.pdf.setTitle(self.prs.core_properties.title)
        self.pdf.setAuthor('Leadprocure')
        self.pages, self.issues = [], []

    def begin(self, n, title, dark=False, section=None):
        self.n, self.dark = n, dark
        self.fg, self.bg = (WHITE, BLACK) if dark else (BLACK, WHITE)
        self.secondary = LIGHT if dark else MUTED
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.page = {'slide':n, 'title':title, 'theme':'dark' if dark else 'light', 'text':[], 'images':[], 'elements':[],
                     'speaker_notes':CONTENT[n-1]['speakerNotes'], 'sources':CONTENT[n-1]['sources']}
        self.pages.append(self.page)
        self.rect(0,0,W,H,self.bg,kind='background')
        if section: self.text(az_upper(section),72,45,1100,15,self.secondary,font='BodyBold',tracking=0.5)

    def text(self, s, x, y, w, size=26, color=None, font='Body', leading=1.2, align='left', max_lines=None, tracking=0):
        color = color or self.fg
        lines=[]
        for p in s.split('\n'):
            current=''
            for word in p.split():
                candidate = (current+' '+word).strip()
                if current and pdfmetrics.stringWidth(candidate,font,size)>w:
                    lines.append(current); current=word
                else: current=candidate
            lines.append(current)
        if max_lines is not None and len(lines)>max_lines:
            self.issues.append(f'Slide {self.n}: {len(lines)} lines > {max_lines}: {s}')
        ascent, descent = pdfmetrics.getAscentDescent(font,size)
        h = (len(lines)-1)*size*leading + ascent-descent
        if x<0 or y<0 or x+w>W+1 or y+h>H-8:
            self.issues.append(f'Slide {self.n}: text outside canvas ({x},{y},{w},{h}): {s}')
        family = 'Barlow Condensed' if font.startswith('Display') else 'IBM Plex Sans'
        for i, l in enumerate(lines):
            measured = pdfmetrics.stringWidth(l,font,size)
            if measured>w+1: self.issues.append(f'Slide {self.n}: overlong word: {l}')
            yy = y+i*size*leading
            # A separate native text box per line prevents office engines from rewrapping prose.
            box=self.slide.shapes.add_textbox(Pt(x),Pt(yy),Pt(w+1),Pt(size*1.55))
            tf=box.text_frame; tf.clear(); tf.word_wrap=False
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            tf.vertical_anchor=MSO_ANCHOR.TOP
            p=tf.paragraphs[0]; p.text=l
            p.font.name=family; p.font.size=Pt(size); p.font.bold=font in ('Display','BodyBold')
            p.font.color.rgb=rgb(color); p.space_after=Pt(0); p.space_before=Pt(0)
            p.line_spacing=1.0
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            # ReportLab and native PPTX share text, bounds, font files and point sizes.
            self.pdf.setFont(font,size); self.pdf.setFillColor(HexColor('#'+color))
            baseline=H-yy-ascent
            if align=='center': self.pdf.drawCentredString(x+w/2,baseline,l)
            elif align=='right': self.pdf.drawRightString(x+w,baseline,l)
            else: self.pdf.drawString(x,baseline,l)
        self.page['text'].append(s)
        self.page['elements'].append({'type':'text','text':s,'bounds':[x,y,w,h],'font':font,'font_size':size,'color':color})
        return h

    def rect(self,x,y,w,h,fill=None,stroke=None,width=1,kind='shape'):
        sh=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
        if fill: sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
        else: sh.fill.background()
        if stroke: sh.line.color.rgb=rgb(stroke); sh.line.width=Pt(width)
        else: sh.line.fill.background()
        self.pdf.setLineWidth(width)
        self.pdf.setFillColor(HexColor('#'+(fill or self.bg)))
        self.pdf.setStrokeColor(HexColor('#'+(stroke or self.bg)))
        self.pdf.rect(x,H-y-h,w,h,stroke=int(bool(stroke)),fill=int(bool(fill)))
        self.page['elements'].append({'type':kind,'bounds':[x,y,w,h]})
        return sh

    def ellipse(self,x,y,w,h,fill=None,stroke=None,width=2):
        sh=self.slide.shapes.add_shape(MSO_SHAPE.OVAL,Pt(x),Pt(y),Pt(w),Pt(h))
        if fill: sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
        else: sh.fill.background()
        if stroke: sh.line.color.rgb=rgb(stroke); sh.line.width=Pt(width)
        else: sh.line.fill.background()
        self.pdf.setFillColor(HexColor('#'+(fill or self.bg)))
        self.pdf.setStrokeColor(HexColor('#'+(stroke or self.bg))); self.pdf.setLineWidth(width)
        self.pdf.ellipse(x,H-y-h,x+w,H-y,stroke=int(bool(stroke)),fill=int(bool(fill)))

    def line(self,x1,y1,x2,y2,color=None,width=2,dash=False):
        color=color or self.fg
        sh=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2))
        sh.line.color.rgb=rgb(color); sh.line.width=Pt(width)
        if dash: sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        self.pdf.setStrokeColor(HexColor('#'+color)); self.pdf.setLineWidth(width)
        self.pdf.setDash([10,9] if dash else [])
        self.pdf.line(x1,H-y1,x2,H-y2); self.pdf.setDash([])

    def arrow(self,x1,y1,x2,y2,color=None,width=2,dash=False,head=10):
        self.line(x1,y1,x2,y2,color,width,dash)
        a=math.atan2(y2-y1,x2-x1)
        for da in (-.48,.48):
            self.line(x2,y2,x2-head*math.cos(a+da),y2-head*math.sin(a+da),color,width)

    def image(self,path,x,y,w,h):
        self.slide.shapes.add_picture(str(path),Pt(x),Pt(y),width=Pt(w),height=Pt(h))
        self.pdf.drawImage(str(path),x,H-y-h,width=w,height=h)
        self.page['images'].append({'file':str(path.relative_to(ROOT)), 'bounds':[x,y,w,h]})

    def person(self,cx,y,scale=1,color=None):
        color=color or self.fg
        self.ellipse(cx-23*scale,y,46*scale,46*scale,stroke=color,width=2.5)
        self.line(cx-45*scale,y+102*scale,cx-45*scale,y+79*scale,color,2.5)
        self.line(cx-45*scale,y+79*scale,cx-22*scale,y+62*scale,color,2.5)
        self.line(cx-22*scale,y+62*scale,cx+22*scale,y+62*scale,color,2.5)
        self.line(cx+22*scale,y+62*scale,cx+45*scale,y+79*scale,color,2.5)
        self.line(cx+45*scale,y+79*scale,cx+45*scale,y+102*scale,color,2.5)
        self.line(cx-45*scale,y+102*scale,cx+45*scale,y+102*scale,color,2.5)

    def document(self,x,y,w=90,h=116,lines=3,locked=False):
        self.rect(x,y,w,h,stroke=self.fg,width=2.3)
        for i in range(lines): self.line(x+18,y+28+i*23,x+w-18,y+28+i*23,self.secondary,2)
        if locked:
            self.rect(x+w-25,y+h-36,40,32,self.bg,self.fg,2)
            self.ellipse(x+w-16,y+h-55,22,30,stroke=self.fg,width=2)
            self.rect(x+w-24,y+h-34,38,29,self.bg,self.fg,2)

    def title(self,title,x=72,y=91,w=1136,size=68,max_lines=2):
        return self.text(az_upper(title),x,y,w,size,font='Display',leading=1.02,max_lines=max_lines)

    def end(self,footer=True):
        if footer:
            self.text('Leadprocure',72,676,250,17,font='BodyBold',color=self.secondary)
            self.text('KORPORATİV SATINALMA',856,679,300,12,color=self.secondary,align='right')
        self.text(f'{self.n:02d}',1170,676,38,17,color=self.secondary,align='right')
        self.slide.notes_slide.notes_text_frame.text=self.page['speaker_notes']+'\n\nMənbələr:\n'+'\n'.join(self.page['sources'])
        self.pdf.showPage()

    def save(self):
        self.pdf.save()
        self.prs.save(OUT/'Leadprocure_Korporativ_Teqdimat_AZ_v2.pptx')
        (OUT/'slide-manifest.json').write_text(json.dumps(self.pages,ensure_ascii=False,indent=2))
        (OUT/'qa-presentation.json').write_text(json.dumps({'slide_count':len(self.pages),'text_layout_issues':self.issues,
            'screenshots':0,'cover_illustrations':sum(len(p['images']) for p in self.pages),
            'design':'C / Monograph','fonts':FONT_FILES,'native_pptx_render':'pending'},ensure_ascii=False,indent=2))
        notes=['# Leadprocure — Məruzəçi qeydləri\n','13 slayd · Monograph · Azərbaycan dili\n']
        for p in self.pages:
            notes += [f"## {p['slide']:02d}. {p['title']}\n",p['speaker_notes']+'\n','Mənbələr: '+ '; '.join(p['sources'])+'\n']
        (OUT/'Meruzeci_Qeydleri_AZ.md').write_text('\n'.join(notes))
        doc=pymupdf.open(OUT/'Leadprocure_Korporativ_Teqdimat_AZ_v2.pdf')
        for i,p in enumerate(doc):
            p.get_pixmap(matrix=pymupdf.Matrix(1,1),alpha=False).save(PREVIEW/f'slide-{i+1:02d}.png')
        # Small sequential renders and contact sheets, never full browser E2E.
        for start in (0,6,12):
            subset=list(range(start,min(start+6,len(doc))))
            board=Image.new('RGB',(1280,math.ceil(len(subset)/2)*392),'#E8E8E8')
            d=ImageDraw.Draw(board); font=ImageFont.truetype(str(ROOT/'fonts'/'IBMPlexSans-Regular.ttf'),15)
            for j,i in enumerate(subset):
                im=Image.open(PREVIEW/f'slide-{i+1:02d}.png').convert('RGB'); im.thumbnail((620,349))
                x=10+(j%2)*640;y=10+(j//2)*392
                board.paste(im,(x,y)); d.text((x,y+356),f'{i+1:02d} / Leadprocure',font=font,fill='#333333')
            board.save(OUT/f'contact-sheet-{start//6+1}.jpg',quality=92)
        assert len(doc)==13
        assert len(self.prs.slides)==13
        print(json.dumps({'slides':len(doc),'layout_issues':self.issues},ensure_ascii=False))
        if self.issues: raise SystemExit(1)

d=Deck()

# 01 — Preserve the selected cover's monumental type and sculptural network.
d.begin(1,'Satınalmanı strateji üstünlüyə çevirin.',dark=True)
d.image(ROOT/'assets'/'monograph-network.png',0,0,W,H)
d.text('Leadprocure',72,74,530,82,font='Display',leading=1)
d.text('Satınalmanı\nstrateji üstünlüyə\nçevirin.',72,294,540,78,font='Display',leading=1.02,max_lines=3)
d.text('Korporativ satınalma\nplatforması',72,591,500,25,LIGHT,leading=1.25)
d.end(footer=False)

# 02 — The selected white value page, with more legible supporting labels.
d.begin(2,'Hər qərar görünür. Hər mərhələ nəzarətdədir.',section='İdarəetmə dəyəri')
d.title('Hər qərar görünür.\nHər mərhələ nəzarətdədir.',size=68)
values=[('01','Xərc şəffaflığı','Büdcədən faktiki xərcə'),('02','Rəqabətli seçim','Qaydalı və əsaslandırılmış qərar'),('03','İcra intizamı','Sifarişdən malın qəbuluna')]
for x,(n,t,b) in zip((72,478,884),values):
    d.text(n,x,290,260,180,SILVER,font='Display',leading=1)
    d.text(az_upper(t),x,499,330,38,font='Display')
    d.text(b,x,556,320,23,MUTED,max_lines=2)
d.arrow(289,411,433,411,SILVER,2);d.arrow(695,411,839,411,SILVER,2)
d.end()

# 03 — Separation of duties and locked evaluation criteria.
d.begin(3,'Satınalma qaydaları qərardan əvvəl müəyyən edilir',section='Daxili nəzarət')
d.title('Qərardan əvvəl\nqaydalar.',w=450,size=64)
d.text('Səlahiyyətlər ayrılır.\nMeyarlar əvvəlcədən\nmüəyyən edilir.',72,335,380,27,MUTED,leading=1.4)
for cx,y in [(580,252),(820,252)]:d.person(cx,y,1)
d.document(1040,250,92,110,3,True)
d.arrow(638,305,757,305,MID,2);d.arrow(882,305,1017,305,MID,2)
d.text('Tələb\nmüəllifi',490,395,180,34,font='Display',align='center')
d.text('Səlahiyyətli\ntəsdiqləyən',710,395,220,34,font='Display',align='center')
d.text('Sabit\nmeyarlar',1005,395,162,34,font='Display',align='center')
d.line(490,498,1167,498,HAIR,1)
d.text('Müəllif öz tələbini təsdiqləmir.',490,522,710,23,MUTED)
d.text('Məbləğə uyğun təsdiq zənciri.',72,617,1136,26,font='BodyBold')
d.end()

# 04 — A branching diagram explains conditional follow-up behavior.
d.begin(4,'Rəqabət təchizatçının prosesə qoşulmasından başlayır',dark=True,section='Təchizatçı iştirakı')
d.title('Dəvətdən\ncavaba qədər.',size=72)
d.text('Dəvət',72,379,190,44,font='Display')
d.line(237,405,312,405,LIGHT,2)
d.line(312,323,312,510,LIGHT,2)
d.arrow(312,323,420,323,LIGHT,2);d.arrow(312,510,420,510,LIGHT,2)
d.ellipse(444,308,18,18,fill=WHITE)
d.text('Təklif təqdim edilib',486,298,360,32,font='Display')
d.arrow(848,323,906,323,LIGHT,2)
d.text('Xatırlatma\ndayanır',936,288,263,34,font='Display')
d.ellipse(444,496,18,18,stroke=SILVER,width=2)
d.text('Cavab gözlənilir',486,485,345,32,font='Display')
d.arrow(848,510,906,510,LIGHT,2)
d.text('Planlı\nxatırlatma',936,475,263,34,font='Display')
d.text('Sual-cavab məlumatı uyğun iştirakçılar üçün ortaqdır.',72,616,1100,23,LIGHT)
d.end()

# 05 — Two clearly separated states, no fake interface or readable sealed price.
d.begin(5,'Təkliflər açılanadək qiymətlər məxfidir',section='Məxfi RFQ / RFP')
d.title('Təkliflər açılanadək\nqiymətlər məxfidir.',size=68)
for x in (120,252,384):d.document(x,316,84,120,0,True);d.text('•••',x+12,354,60,31,font='BodyBold',align='center')
d.text('Açılışdan əvvəl',104,496,410,36,font='Display')
d.text('Alıcı məbləği görmür.',104,554,430,25,MUTED)
d.arrow(518,377,649,377,MID,2)
d.ellipse(659,319,116,116,stroke=BLACK,width=2)
d.text('AÇILIŞ',669,360,96,24,font='Display',align='center')
d.arrow(785,377,862,377,MID,2)
for y in (320,365,410):
    d.line(899,y,1180,y,BLACK,2);d.ellipse(884,y-5,10,10,fill=BLACK)
d.text('Qiymətləndirmə',889,496,330,36,font='Display')
d.text('Əvvəlcədən müəyyən\nedilmiş meyarlarla.',889,554,319,24,MUTED,max_lines=2)
d.end()

# 06 — The selected auction centerpiece. Final step is an allowed ceiling, not a bid.
d.begin(6,'Revers-auksion: rəqabət real vaxtda.',dark=True,section='Qiymət rəqabəti')
d.title('Revers-auksion:\nrəqabət real vaxtda.',size=68)
for x,y,amount in [(72,256,'32 000'),(469,351,'31 200'),(866,446,'30 880')]:
    d.text(amount,x,y,320,108,font='Display',leading=1,max_lines=1)
    d.text('AZN',x+280,y+65,72,21,LIGHT,font='BodyBold')
d.line(72,381,442,381,WHITE,3);d.line(442,381,442,476,WHITE,3)
d.line(442,476,838,476,WHITE,3);d.line(838,476,838,571,WHITE,3)
d.line(838,571,1207,571,WHITE,3,dash=True)
d.text('Start',72,397,300,22,LIGHT)
d.text('Ən yaxşı təklif',469,492,330,22,LIGHT)
d.text('Növbəti maksimum təklif',866,588,342,21,LIGHT,max_lines=1)
d.text('Minimum addım',72,641,290,22,LIGHT)
d.text('Son saniyədə uzatma',440,641,350,22,LIGHT)
d.text('Şəffaf qaydalar',928,641,280,22,LIGHT,align='right')
d.text('İllüstrativ nümunə',906,45,302,15,LIGHT,align='right')
d.end()

# 07 — Explicit RFQ demonstration, not claimed customer ROI or an auction result.
d.begin(7,'Seçimi rəqəmlərlə görünən edin',section='RFQ / demo nümunəsi')
d.title('Seçimi rəqəmlərlə\ngörünən edin.',size=68)
d.text('İlkin smeta',72,291,500,22,MUTED)
d.text('57 000',72,323,570,100,font='Display')
d.text('AZN',399,387,85,23,MUTED,font='BodyBold')
d.rect(72,432,720,10,LIGHT)
d.text('Seçilən təklif',72,468,500,22,MUTED)
d.text('51 500',72,501,570,100,font='Display')
d.text('AZN',399,565,85,23,MUTED,font='BodyBold')
d.rect(72,610,720*(51500/57000),10,BLACK)
d.line(847,294,847,620,HAIR,1)
d.text('9,65%',904,313,304,100,font='Display',max_lines=1)
d.text('Smetaya nisbətən fərq',904,433,304,24,MUTED,max_lines=2)
d.text('5 500 AZN',904,515,304,53,font='Display',max_lines=1)
d.text('Demo məlumatları · ƏDV-siz · real müştəri nəticəsi deyil.',72,641,1136,18,MUTED)
d.end()

# 08 — Explainable decision support, human decision remains distinct.
d.begin(8,'AI tövsiyəsi. Qərar sizdədir.',dark=True,section='Qərar dəstəyi')
d.title('AI tövsiyəsi.\nQərar sizdədir.',size=72)
for y,t in [(315,'Qiymət'),(438,'Çatdırılma müddəti'),(561,'Performans tarixçəsi')]:
    d.ellipse(73,y+7,14,14,fill=SILVER)
    d.text(t,106,y-2,340,29,font='BodyBold',max_lines=1)
    d.line(430,y+19,540,438,LIGHT,2)
d.arrow(540,438,605,438,LIGHT,2)
d.ellipse(617,334,208,208,stroke=WHITE,width=2)
d.text('İzah edilən\ntövsiyə',638,393,168,32,font='Display',align='center')
d.arrow(841,438,948,438,LIGHT,2)
d.person(1050,359,1.05,WHITE)
d.text('Alıcının\nqərarı',941,514,218,35,font='Display',align='center')
d.text('Formal qiymətləndirməni dəstəkləyir; onu əvəz etmir.',72,625,1100,23,LIGHT)
d.end()

# 09 — One contract-to-receipt line, with the supplier's acceptance as a change gate.
d.begin(9,'Razılaşdırılan şərtlər icra boyu izlənir',section='İcra nəzarəti')
d.title('Seçimdən icraya —\nşərtlər izlənir.',size=68)
for x,n,t in [(72,'01','Müqavilə'),(379,'02','Sifariş'),(686,'03','Çatdırılma'),(993,'04','Qəbul')]:
    d.text(n,x,293,210,106,SILVER,font='Display')
    d.text(t,x,425,218,37,font='Display')
    if x<993:d.arrow(x+156,354,x+264,354,MID,2)
d.line(451,482,451,543,MID,2)
d.arrow(451,543,617,543,MID,2)
d.text('Dəyişiklik\ntəklifi',641,510,222,28,font='Display')
d.arrow(837,543,911,543,MID,2)
d.text('Təchizatçının\nrazılığı',935,510,265,28,font='Display')
d.text('Sifariş şərtləri yalnız razılıqdan sonra yenilənir.',72,617,1136,24,MUTED)
d.end()

# 10 — Three-way match as a semantic triangle, not another feature list.
d.begin(10,'Qaimə sifariş və faktiki qəbulla üzləşdirilir',dark=True,section='Qaimə nəzarəti')
d.title('Üç sənəd.\nBir uyğunluq\nyoxlaması.',w=472,size=70,max_lines=3)
d.text('Qiymət və miqdar\nfərqləri görünür.',72,403,390,27,LIGHT,leading=1.35)
d.text('İstisna təsdiqi\nəsaslandırılır.',72,544,390,25,LIGHT,leading=1.35)
d.line(873,274,653,551,SILVER,2);d.line(873,274,1093,551,SILVER,2);d.line(653,551,1093,551,SILVER,2)
for x,y in [(873,274),(653,551),(1093,551)]:d.ellipse(x-9,y-9,18,18,fill=WHITE)
d.text('QAİMƏ',748,204,250,42,font='Display',align='center')
d.text('SİFARİŞ',548,586,210,37,font='Display',align='center')
d.text('FAKTİKİ QƏBUL',968,586,240,37,font='Display',align='center')
d.text('Uyğunluq',738,418,270,43,font='Display',align='center')
d.end()

# 11 — Commitments and recognized spend are deliberately separated.
d.begin(11,'Büdcə öhdəliyi ilə faktiki xərc ayrı görünür',section='Maliyyə baxışı')
d.title('Büdcə öhdəliyi.\nFaktiki xərc.',size=72)
d.line(639,282,639,562,HAIR,1)
for x,n,first,last in [(72,'01','Təsdiqlənmiş tələb','Büdcə öhdəliyi'),(717,'02','Qəbul edilmiş sifariş','Xərc analitikası')]:
    d.text(n,x,286,180,124,SILVER,font='Display')
    d.text(first,x+173,330,316,32,font='Display',max_lines=2)
    d.arrow(x+250,424,x+250,466,MID,2)
    d.text(last,x+173,492,316,35,font='Display',max_lines=1)
d.rect(72,591,1136,61,BLACK)
d.text('Limit aşımı xəbərdarlıq yaradır; təsdiqi bloklamır.',96,609,1088,24,WHITE)
d.end()

# 12 — Executive recap through three questions, with no invented KPIs.
d.begin(12,'Bir proses. Üç rəhbərlik baxışı.',section='Rəhbərlik üçün')
d.title('Bir proses.\nÜç rəhbərlik baxışı.',size=72)
for i,(role,question) in enumerate([('Maliyyə','Vəsait hara yönəlir?'),('Satınalma','Seçim nəyə əsaslanır?'),('Əməliyyat','İcra hansı mərhələdədir?')]):
    y=300+i*112
    d.line(72,y,1208,y,HAIR,1)
    d.text(f'0{i+1}',72,y+30,68,28,MID,font='Display')
    d.text(role,160,y+21,389,49,font='Display')
    d.text(question,623,y+40,585,28,MUTED,max_lines=1)
d.line(72,636,1208,636,HAIR,1)
d.end()

# 13 — Practical next step without invented commercial terms or time promises.
d.begin(13,'Dəyəri bir real satınalma ilə ölçək',dark=True,section='Növbəti addım')
d.title('Dəyəri bir real\nsatınalma ilə ölçək.',size=80)
d.text('Bir kateqoriya. Razılaşdırılmış qaydalar. Aydın meyarlar.',72,295,1136,26,LIGHT,max_lines=1)
for x,n,title,body in [(72,'01','Əhatə','Kateqoriya və rollar'),(478,'02','İcra','Bir satınalma prosesi'),(884,'03','Qiymətləndirmə','Birlikdə ölçülən nəticə')]:
    d.text(n,x,405,220,116,SILVER,font='Display')
    d.text(az_upper(title),x,545,324,42,font='Display',max_lines=1)
    d.text(body,x,605,324,23,LIGHT,max_lines=1)
d.arrow(280,475,430,475,SILVER,2);d.arrow(686,475,836,475,SILVER,2)
d.end()

d.save()
