#!/usr/bin/env python3
"""Render the reviewed two-page Monograph product data sheet; no app/runtime access."""
from __future__ import annotations

import json
from pathlib import Path
import pymupdf as fitz
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas


ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / "datasheet-content.json").read_text())
OUT = ROOT / "output"
PREVIEW = OUT / "preview"
OUT.mkdir(parents=True, exist_ok=True)
PREVIEW.mkdir(parents=True, exist_ok=True)

PDF_PATH = OUT / "Leadprocure_Data_Sheet_AZ.pdf"
MD_PATH = OUT / "Leadprocure_Data_Sheet_AZ.md"
W, H = A4
prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = 'Leadprocure — Məhsul məlumat vərəqi'
prs.core_properties.author = 'Leadprocure'
slide = None
INK = HexColor("#161616")
MID = HexColor("#555555")
SILVER = HexColor("#B8B8B8")
RULE = HexColor("#D6D6D6")
WHITE = HexColor("#FFFFFF")
LEFT = 42.0
RIGHT = W - 42.0
FULL = RIGHT - LEFT
COL_W = (FULL - 30.0) / 2.0
COL_X = [LEFT, LEFT + COL_W + 30.0]

for alias, filename in (
    ("Display", "BarlowCondensed-Bold.ttf"),
    ("DisplayMedium", "BarlowCondensed-SemiBold.ttf"),
    ("Body", "IBMPlexSans-Regular.ttf"),
    ("BodyMedium", "IBMPlexSans-Medium.ttf"),
):
    pdfmetrics.registerFont(TTFont(alias, str(ROOT / "fonts" / filename)))

c = Canvas(str(PDF_PATH), pagesize=A4, pageCompression=1)
c.setTitle("Leadprocure — Məhsul məlumat vərəqi")
c.setAuthor("Leadprocure")
c.setSubject("Korporativ satınalma platforması · Azərbaycan dili")
c.setCreator("Leadprocure Monograph data sheet renderer")

PAGE = 0
BOXES: list[dict] = []
ISSUES: list[str] = []


def az_upper(value: str) -> str:
    return value.replace("i", "İ").replace("ı", "I").upper()


def wrap(value: str, font: str, size: float, width: float) -> list[str]:
    result = []
    for paragraph in value.split("\n"):
        if not paragraph:
            result.append("")
            continue
        line = ""
        for word in paragraph.split():
            trial = (line + " " + word).strip()
            if line and pdfmetrics.stringWidth(trial, font, size) > width:
                result.append(line)
                line = word
            else:
                line = trial
        result.append(line)
    return result


def text(value: str, x: float, top: float, width: float, size: float = 10.25,
         font: str = "Body", color=MID, leading: float | None = None,
         limit: float | None = None, record: bool = True) -> float:
    leading = leading if leading is not None else size * 1.28
    lines = wrap(value, font, size, width)
    height = len(lines) * leading
    if limit is not None and height > limit + 0.1:
        ISSUES.append(f"Page {PAGE}: text exceeds allocation by {height-limit:.2f}pt: {value[:80]}")
    if top + height > H - 20:
        ISSUES.append(f"Page {PAGE}: text exceeds page bottom: {value[:80]}")
    c.setFillColor(color)
    c.setFont(font, size)
    for i, line_value in enumerate(lines):
        line_width = pdfmetrics.stringWidth(line_value, font, size)
        if line_width > width + 0.1:
            ISSUES.append(f"Page {PAGE}: unwrapped word exceeds box: {line_value}")
        c.drawString(x, H - top - size - i * leading, line_value)
        ascent = pdfmetrics.getAscent(font, size)
        box = slide.shapes.add_textbox(Pt(x), Pt(top + size - ascent + i * leading), Pt(width + 1), Pt(size * 1.55))
        tf = box.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = line_value
        p.font.name = 'Barlow Condensed' if font.startswith('Display') else 'IBM Plex Sans'
        p.font.bold = font.startswith('Display')
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(*(round(v * 255) for v in (color.red, color.green, color.blue)))
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.0
    if record:
        BOXES.append({"page": PAGE, "text": value, "x": x, "top": top,
                      "width": width, "height": height, "font": font, "fontSize": size})
    return height


def rule(x: float, top: float, width: float, color=RULE, thickness: float = 0.5):
    c.setStrokeColor(color)
    c.setLineWidth(thickness)
    c.line(x, H - top, x + width, H - top)
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x), Pt(top), Pt(x + width), Pt(top))
    sh.line.width = Pt(thickness)
    sh.line.color.rgb = RGBColor(*(round(v * 255) for v in (color.red, color.green, color.blue)))
    sh._element.spPr.append(OxmlElement('a:effectLst'))


def panel(x: float, top: float, width: float, height: float, color=INK):
    c.setFillColor(color)
    c.rect(x, H - top - height, width, height, fill=1, stroke=0)
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(top), Pt(width), Pt(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(*(round(v * 255) for v in (color.red, color.green, color.blue)))
    sh.line.fill.background()
    sh._element.spPr.append(OxmlElement('a:effectLst'))


def footer():
    rule(LEFT, 804, FULL)
    text("LEADPROCURE  /  MƏHSUL MƏLUMAT VƏRƏQİ  /  AZ", LEFT, 811, FULL - 60,
         size=8.5, font="BodyMedium", color=MID, leading=10.5)
    text(f"{PAGE:02} / 02", RIGHT - 38, 811, 38, size=8.5,
         font="BodyMedium", color=INK, leading=10.5)


def section(title: str, x: float, top: float, width: float) -> float:
    height = text(title, x, top, width, size=19, font="Display", color=INK, leading=22)
    rule(x, top + height + 6, width)
    return top + height + 17


def labeled(title: str, body: str, x: float, top: float, width: float,
            size: float = 10.25) -> float:
    h = text(title, x, top, width, size=size, font="BodyMedium", color=INK, leading=13.2)
    h += 4
    h += text(body, x, top + h, width, size=size, font="Body", leading=13.2)
    return h


# PAGE 1 — white body, dark brand band, ten numbered capabilities.
PAGE = 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
panel(0, 0, W, 130)
text(DATA["title"], LEFT, 18, 330, size=46, font="Display", color=WHITE, leading=51)
text(az_upper(DATA["subtitle"]), RIGHT - 180, 28, 180, size=9.5,
     font="BodyMedium", color=SILVER, leading=12.3)
text(DATA["tagline"], LEFT, 83, FULL, size=20, font="DisplayMedium",
     color=WHITE, leading=23, limit=46)

intro_bottom = 147 + text(DATA["intro"], LEFT, 147, FULL, size=10.25, color=INK, leading=13.2)
target_top = intro_bottom + 8
target_bottom = target_top + text(DATA["targetBuyer"], LEFT, target_top, FULL, size=10.25,
                                  color=MID, leading=13.2)
section_top = target_bottom + 12
grid_top = section("Əsas imkanlar", LEFT, section_top, FULL)
row_h = (793.0 - grid_top) / 5.0

for index, item in enumerate(DATA["capabilities"]):
    col, row = divmod(index, 5)
    x = COL_X[col]
    y = grid_top + row * row_h
    text(f"{index+1:02}", x - 19, y + 3, 16, size=10.5, font="DisplayMedium", color=MID, leading=13)
    title_h = text(item["title"], x, y, COL_W, size=17, font="Display", color=INK, leading=19.2)
    body_h = text(item["description"], x, y + title_h + 5, COL_W, size=10.25,
                  font="Body", color=MID, leading=13.0)
    required = title_h + 5 + body_h
    if required > row_h - 7:
        ISSUES.append(f"Page 1 capability {index+1}: only {row_h-required:.2f}pt remaining in row")

footer()
c.showPage()

# PAGE 2 — specification columns, shared condition notes, restrained closing band.
PAGE = 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
panel(0, 0, W, 112)
text("Leadprocure", LEFT, 20, 230, size=17, font="Display", color=WHITE, leading=21)
text("Profil və tətbiq şərtləri", LEFT, 48, FULL, size=43, font="Display", color=WHITE, leading=49)

x1, x2 = COL_X
y = section("İş mühiti", x1, 133, COL_W)
y += labeled("Rollar", " · ".join(DATA["roles"]), x1, y, COL_W) + 11
y += labeled("Dillər", " · ".join(DATA["languages"]), x1, y, COL_W) + 11
y += labeled("Maliyyə məlumatları", " ".join(DATA["financialFields"]), x1, y, COL_W)
environment_bottom = y

y = section("Yerləşdirmə", x2, 133, COL_W)
for item in DATA["deployment"]:
    y += labeled(item["title"], item["description"], x2, y, COL_W) + 12
deployment_bottom = y - 12

second_row = max(environment_bottom, deployment_bottom) + 21
y = section("Nəzarət mexanizmləri", x1, second_row, COL_W)
for value in DATA["security"]:
    panel(x1, y + 5, 2.5, 2.5, MID)
    y += text(value, x1 + 10, y, COL_W - 10, size=10.25, leading=13.2) + 7
security_bottom = y - 7

y = section("Məlumat mübadiləsi", x2, second_row, COL_W)
for item in DATA["integrations"]:
    y += labeled(item["title"], item["description"], x2, y, COL_W) + 11
integration_bottom = y - 11

conditions_top = max(security_bottom, integration_bottom) + 20
y = section("Funksiyaların tətbiq şərtləri", LEFT, conditions_top, FULL)
for value in DATA["caveats"]:
    panel(LEFT, y + 4.5, 2.2, 2.2, MID)
    y += text(value, LEFT + 10, y, FULL - 10, size=9.0, leading=11.6) + 4
conditions_bottom = y - 4

pilot_top = max(conditions_bottom + 15, 730)
pilot_h = 62
if pilot_top + pilot_h > 795:
    ISSUES.append(f"Page 2 pilot panel reaches {pilot_top+pilot_h:.1f}pt; target bottom 795pt")
panel(LEFT, pilot_top, FULL, pilot_h)
text("Pilot satınalma ilə başlayın", LEFT + 13, pilot_top + 8, FULL - 26,
     size=18, font="Display", color=WHITE, leading=21)
text(DATA["pilot"], LEFT + 13, pilot_top + 32, FULL - 26,
     size=10.25, font="Body", color=WHITE, leading=13.2, limit=27)

footer()
c.showPage()
c.save()
prs.save(OUT / 'Leadprocure_Data_Sheet_AZ.pptx')


def markdown() -> str:
    lines = [f"# {DATA['title']} — {DATA['subtitle']}", "", DATA["tagline"], "",
             "## Səhifə 1 — Məhsul və əsas imkanlar", "", DATA["intro"], "", DATA["targetBuyer"], ""]
    for i, item in enumerate(DATA["capabilities"], 1):
        lines += [f"### {i:02}. {item['title']}", "", item["description"], ""]
    lines += ["## Səhifə 2 — Profil və tətbiq şərtləri", ""]
    for title, key in (("Rollar", "roles"), ("Dillər", "languages"), ("Maliyyə məlumatları", "financialFields")):
        lines += [f"### {title}", ""] + [f"- {v}" for v in DATA[key]] + [""]
    for title, key in (("Yerləşdirmə", "deployment"), ("Məlumat mübadiləsi", "integrations")):
        lines += [f"### {title}", ""]
        for item in DATA[key]:
            lines += [f"**{item['title']}.** {item['description']}", ""]
    for title, key in (("Nəzarət mexanizmləri", "security"), ("Funksiyaların tətbiq şərtləri", "caveats")):
        lines += [f"### {title}", ""] + [f"- {v}" for v in DATA[key]] + [""]
    lines += ["### Pilot", "", DATA["pilot"], ""]
    return "\n".join(lines)


MD_PATH.write_text(markdown())

doc = fitz.open(PDF_PATH)
pdf_text = "\n".join(page.get_text() for page in doc)
normalized_text = " ".join(pdf_text.split())
missing = []
for box in BOXES:
    if " ".join(box["text"].split()) not in normalized_text:
        missing.append(box["text"])
for i, page in enumerate(doc, 1):
    page.get_pixmap(matrix=fitz.Matrix(1.55, 1.55), alpha=False).save(PREVIEW / f"datasheet-{i:02}.png")

qa = {
    "document": PDF_PATH.name,
    "editablePowerPoint": "Leadprocure_Data_Sheet_AZ.pptx",
    "powerPointPages": len(prs.slides),
    "pageCount": len(doc),
    "requiredPageCount": 2,
    "pageSize": "A4 portrait",
    "pageSizesPt": [[round(p.rect.width, 2), round(p.rect.height, 2)] for p in doc],
    "displayFont": "Barlow Condensed",
    "bodyFont": "IBM Plex Sans",
    "minimumBodyFontPt": 10.25,
    "minimumConditionFontPt": 9.0,
    "minimumFooterFontPt": 8.5,
    "capabilities": len(DATA["capabilities"]),
    "layoutIssues": ISSUES,
    "missingText": missing,
    "replacementCharacterPresent": "\ufffd" in pdf_text,
    "sourceCopyPreserved": not missing,
    "page1GridTopPt": round(grid_top, 2),
    "page1CapabilityRowHeightPt": round(row_h, 2),
    "page2ConditionsBottomPt": round(conditions_bottom, 2),
    "page2PilotBottomPt": round(pilot_top + pilot_h, 2),
    "visualReview": "PENDING",
    "docx": "NOT RUN: python-docx is unavailable; no dependency installation requested. Editable Markdown source supplied.",
    "pdfBytes": PDF_PATH.stat().st_size,
}
(OUT / "qa-datasheet.json").write_text(json.dumps(qa, ensure_ascii=False, indent=2) + "\n")
if len(doc) != 2:
    raise SystemExit(f"Expected exactly 2 pages; got {len(doc)}")
if missing or ISSUES:
    print(json.dumps(qa, ensure_ascii=False, indent=2))
    raise SystemExit("Data sheet requires layout review")
print(json.dumps(qa, ensure_ascii=False, indent=2))
