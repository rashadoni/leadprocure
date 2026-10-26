#!/usr/bin/env python3
"""Build editable PPTX + matching PDF from audited, local demo screenshots.

Run with python-pptx, Pillow, reportlab and pymupdf on PYTHONPATH.
No network, app build, runtime access or production changes.
"""
from pathlib import Path
import json, math
from io import BytesIO
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
import fitz

ROOT=Path(__file__).resolve().parent
ASSETS=ROOT/'assets'
OUT=ROOT/'output'
OUT.mkdir(exist_ok=True)
W,H=1152,648
BG='F8F7F3'; INK='202B27'; MUTED='607069'; GREEN='087F62'; FOREST='164B3C'
MINT='E6F1E9'; LINE='DCE4DE'; WHITE='FFFFFF'; GOLD='B38C45'; CLAY='AA583E'
FONT='Lato'
pdfmetrics.registerFont(TTFont('Lato','/usr/share/fonts/truetype/lato/Lato-Regular.ttf'))
pdfmetrics.registerFont(TTFont('Lato-Bold','/usr/share/fonts/truetype/lato/Lato-Bold.ttf'))
prs=Presentation(); prs.slide_width=Pt(W); prs.slide_height=Pt(H)
prs.core_properties.title='Leadprocure — Korporativ satınalma platforması'
prs.core_properties.subject='Potensial korporativ müştərilər üçün məhsul təqdimatı'
prs.core_properties.author='Leadprocure'
pdf=canvas.Canvas(str(OUT/'Leadprocure_Korporativ_Teqdimat_AZ.pdf'),pagesize=(W,H))
pdf.setTitle(prs.core_properties.title); pdf.setAuthor('Leadprocure')
manifest=[]; violations=[]; all_sources=set()

def color(h): return RGBColor.from_string(h)
def rect(x,y,w,h,fill,stroke=None,r=0):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color(fill)
    sh.line.fill.background() if not stroke else None
    if stroke: sh.line.color.rgb=color(stroke); sh.line.width=Pt(.75)
    if r: sh.adjustments[0]=min(.15,r/min(w,h))
    pdf.setFillColor(HexColor('#'+fill)); pdf.setStrokeColor(HexColor('#'+(stroke or fill)))
    if r: pdf.roundRect(x,H-y-h,w,h,r,fill=1,stroke=int(bool(stroke)))
    else: pdf.rect(x,H-y-h,w,h,fill=1,stroke=int(bool(stroke)))
    return sh

def wrap(s,w,size,bold=False):
    lines=[]; font='Lato-Bold' if bold else 'Lato'
    for p in s.split('\n'):
        line=''
        for word in p.split():
            cand=(line+' '+word).strip()
            if line and pdfmetrics.stringWidth(cand,font,size)>w: lines.append(line); line=word
            else: line=cand
        lines.append(line)
    return lines

def text(s,x,y,w,size=20,fill=INK,bold=False,max_h=None,leading=1.22,align='left'):
    lines=wrap(s,w,size,bold); h=len(lines)*size*leading+5
    if max_h and h>max_h: violations.append(f'slide {len(manifest)} text height {h:.1f}>{max_h}: {s[:60]}')
    if y+h>H-8: violations.append(f'slide {len(manifest)} text beyond page: {s[:60]}')
    box=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    from pptx.enum.text import PP_ALIGN
    for i,l in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=l
        p.font.name=FONT; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color(fill)
        p.space_after=Pt(0); p.space_before=Pt(0); p.line_spacing=Pt(size*leading)
        if align=='center': p.alignment=PP_ALIGN.CENTER
        if align=='right': p.alignment=PP_ALIGN.RIGHT
        pdf.setFont('Lato-Bold' if bold else 'Lato',size); pdf.setFillColor(HexColor('#'+fill))
        baseline=H-y-size*.91-i*size*leading
        if align=='center': pdf.drawCentredString(x+w/2,baseline,l)
        elif align=='right': pdf.drawRightString(x+w,baseline,l)
        else: pdf.drawString(x,baseline,l)
    manifest[-1]['text'].append(s)
    return h

def line(x1,y1,x2,y2,c=LINE,width=1):
    from pptx.enum.shapes import MSO_CONNECTOR
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2))
    sh.line.color.rgb=color(c); sh.line.width=Pt(width)
    pdf.setStrokeColor(HexColor('#'+c)); pdf.setLineWidth(width); pdf.line(x1,H-y1,x2,H-y2)

def shot(name,x,y,w,h,crop=None,caption=True):
    path=ASSETS/name; im=Image.open(path).convert('RGB')
    if crop: im=im.crop(crop)
    iw,ih=im.size; scale=min(w/iw,h/ih); dw,dh=iw*scale,ih*scale
    xx=x+(w-dw)/2; yy=y+(h-dh)/2
    rect(xx-5,yy-5,dw+10,dh+10,WHITE,LINE,r=6)
    stream=BytesIO(); im.save(stream,format='PNG'); stream.seek(0)
    slide.shapes.add_picture(stream,Pt(xx),Pt(yy),width=Pt(dw),height=Pt(dh))
    pdf.drawImage(ImageReader(im),xx,H-yy-dh,width=dw,height=dh)
    manifest[-1]['images'].append({'file':name,'crop':crop,'rect':[xx,yy,dw,dh]})
    if caption: text('AZ demo interfeysi'+(' · ekran fraqmenti' if crop else ''),xx,yy+dh+10,dw,10.5,MUTED)

def start(kicker,title,sub='',notes='',sources=(),dark=False):
    global slide
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    manifest.append({'slide':len(manifest)+1,'title':title,'text':[],'images':[],'notes':notes,'sources':list(sources)})
    rect(0,0,W,H,FOREST if dark else BG)
    text('Leadprocure',46,25,200,17,WHITE if dark else GREEN,True)
    text(kicker.replace('i','İ').replace('ı','I').upper(),750,29,356,10.5,'BDDED0' if dark else MUTED,align='right')
    line(46,60,1106,60,'41705F' if dark else LINE)
    text(title,46,85,1060,35,WHITE if dark else INK,True,max_h=100)
    if sub: text(sub,46,137,1060,17,'C7E2D5' if dark else MUTED,max_h=65)
    text('KORPORATİV SATINALMA  /  MƏHSUL TƏQDİMATI',46,617,800,9.5,'BDDED0' if dark else MUTED)
    text(f'{len(manifest):02}',1056,613,50,13,'BDDED0' if dark else GREEN,True,align='right')
    all_sources.update(sources)

def end():
    m=manifest[-1]
    slide.notes_slide.notes_text_frame.text=m['notes']+'\n\nMənbələr:\n'+'\n'.join(m['sources'])+'\nEkran görüntüləri: saxlanmış AZ demo video bələdçiləri. Ekrandakı rəqəmlər və tarixlər nümayiş məlumatlarıdır.'
    pdf.showPage()

def card(n,title,body,x,y,w=330,h=205,fill=WHITE):
    rect(x,y,w,h,fill,LINE,r=8)
    text(n,x+22,y+18,w-44,13,GREEN,True)
    text(title,x+22,y+48,w-44,23,INK,True,max_h=65)
    text(body,x+22,y+112,w-44,17,MUTED,max_h=h-120)

def feature(n,title,body,x,y,w=300):
    text(n,x,y,35,13,GREEN,True)
    text(title,x+44,y-2,w-44,21,INK,True,max_h=58)
    text(body,x+44,y+55,w-44,16,MUTED,max_h=95)

S='apps/api/src/modules/'

# 01 — editorial cover, warm paper rather than a dark generic template.
start('Azərbaycan · müəssisələr üçün','Satınalmanı vahid sistemdə idarə edin',
      'Tələbatdan tenderə, müqavilədən qəbul və ödənişin uçotuna qədər.',
      'Leadprocure korporativ satınalma komandasını, təsdiqləyənləri və təchizatçıları vahid prosesdə birləşdirir. Bu təqdimat funksiyaları real demo ekranları ilə göstərir. Məhsulun təqdimatı, 26 oktyabr 2026.',
      ['README.md:3–33','docs/DEMO.md:23–34'])
rect(46,202,349,349,MINT,r=8)
text('Nəzarət.\nRəqabət.\nİzlənə bilən qərarlar.',70,236,302,34,FOREST,True)
text('Satınalma, maliyyə və rəhbərlik üçün ortaq iş mühiti.',70,420,290,21,MUTED)
shot('buyer-dashboard.az.png',425,211,676,380,caption=False)
text('Məhsulun real AZ demo interfeysi · nümunə məlumatlar',426,587,660,10.5,MUTED)
end()

# 02
start('İdarəetmə dəyəri','Korporativ satınalmanın üç əsas sualına cavab',
      'Nəyə vəsait ayrılır? Qərarı kim verir? Təchizat hansı mərhələdədir?',
      'Faydalar məhsul funksiyalarından çıxarılan əməliyyat dəyəridir; ölçülmüş müştəri ROI-si və ya zəmanətli qənaət təqdim edilmir.',
      [S+'requests/requests.module.ts',S+'dashboard/dashboard.module.ts',S+'contracts/contracts.module.ts'])
card('01 / MALİYYƏ ŞƏFFAFLIĞI','Xərcləri görün','Büdcə, öhdəlik, qəbul edilmiş sifariş və qaimə statusunu əlaqəli izləyin.',46,216,340,255)
card('02 / SƏLAHİYYƏT BÖLGÜSÜ','Qərarları əsaslandırın','Təsdiq zəncirləri, əvvəlcədən müəyyən edilən meyarlar və audit tarixçəsi.',406,216,340,255)
card('03 / ƏMƏLİYYAT NƏZARƏTİ','İcraya nəzarət edin','Təklifdən çatdırılmaya qədər məsul tərəf və növbəti addım görünür.',766,216,340,255)
text('Hədəf: dağınıq yazışmaları və əl ilə izlənən mərhələləri vahid prosesdə toplamaq.',47,520,1045,23,FOREST,True)
end()

# 03
start('Fərqləndirici imkanlar','Sadəcə tender reyestrindən daha geniş imkanlar',
      'Rəqabəti, nəzarəti və təchizatçı ilə əməkdaşlığı birləşdirən funksiyalar.',
      'Bu altı funksiya təqdimatın əsas məhsul fərqləndiriciləridir. Ətraflı nümunələr növbəti slaydlardadır.',
      [S+'sourcing/auction.service.ts','apps/api/src/common/crypto/sealed.service.ts',S+'requests/requests.module.ts',S+'invoices/invoices.module.ts',S+'followups/followups.module.ts'])
items=[('01','Revers-auksion','Qiymət rəqabəti, minimum addım və son saniyə təkliflərinə qarşı uzatma.'),('02','Məxfi təkliflər','Təkliflərin məbləğləri açılışa qədər alıcıdan gizli saxlanılır.'),('03','Çoxmərhələli təsdiq','Məbləğə uyğun razılaşdırma və səlahiyyətlərin ayrılması.'),('04','Üçtərəfli üzləşdirmə','Sifariş, faktiki qəbul və qaimə birlikdə yoxlanılır.'),('05','Avtomatik xatırlatma','Cavab verməmiş dəvətli təchizatçılar üçün izləmə zəncirləri.'),('06','Lotlar üzrə seçim','Bir tenderdə müxtəlif lotları fərqli təchizatçılara vermək imkanı.')]
for i,(n,t,b) in enumerate(items): feature(n,t,b,46+(i%3)*362,218+(i//3)*182,336)
end()

# 04
start('Tam satınalma dövrü','Bir proses. İki tərəf. Aydın məsuliyyət.',
      'Alıcı və təchizatçı eyni əməliyyatın özlərinə aid hissəsində işləyir.',
      'PO — satınalma sifarişidir. Ödəniş mərhələsi sistemdə uçot/status qeydini bildirir, bank köçürməsini deyil.',
      ['docs/DEMO.md:23–34','apps/web/src/main.tsx:46–78'])
steps=[('01','Tələbat','Mövqelər və smeta'),('02','Təsdiq','Rol və məbləğ qaydası'),('03','Rəqabət','Tender və ya auksion'),('04','Seçim','Qiymətləndirmə və qalib'),('05','Müqavilə','Razılaşdırılmış şərtlər'),('06','Sifariş','PO və çatdırılma'),('07','Qəbul','Faktiki miqdar'),('08','Qaimə','Üzləşdirmə və uçot')]
for i,(n,t,b) in enumerate(steps):
    x=46+(i%4)*270; y=215+(i//4)*160
    rect(x,y,250,130,WHITE,LINE,r=8); text(n,x+18,y+15,200,13,GREEN,True); text(t,x+18,y+46,220,23,INK,True); text(b,x+18,y+85,220,15,MUTED)
    if i%4<3: text('→',x+253,y+52,18,20,GREEN)
text('Alıcı: tələbat, təsdiq, seçim və qəbul',47,555,515,18,FOREST,True)
text('Təchizatçı: təklif, çatdırılma və qaimə',586,555,520,18,FOREST,True)
end()

# 05
start('Rəhbərlik və satınalma komandası','Görüləcək işlər bir idarə panelində',
      'Göstəricilər, qərar gözləyən tapşırıqlar və əməliyyat bildirişləri.',
      'Ekranda demo satınalma həcmi, qənaət, aktiv tenderlər, təsdiq gözləyən tələb və büdcə xəbərdarlığı göstərilir.',
      ['apps/web/src/guides.ts:15–43',S+'dashboard/dashboard.module.ts:15–47'])
shot('buyer-dashboard.az.png',46,205,764,390)
feature('01','Prioritetlər görünür','Təsdiq, qalib seçimi və qəbul kimi növbəti addımlar.',842,222,263)
feature('02','Vahid baxış','Aktiv tenderlər, sifarişlər və bildirişlər eyni ekranda.',842,397,263)
end()

# 06
start('Tələbatın idarə edilməsi','Satınalma dəqiq tələbatdan başlayır',
      'Kateqoriya, miqdar və qiymət smetası sonrakı qərarların əsasını yaradır.',
      'İnterfeysdə daxili satınalma tələbi “Sifarişlər”, purchase order isə “Sifarişlər (PO)” kimi göstərilir. Təqdimatda qarışıqlığı azaltmaq üçün “tələbat” və “PO” ayrılır.',
      [S+'requests/requests.module.ts:170–227','scripts/demo-fill.mjs:197–201'])
shot('buyer-requests.az.png',46,205,727,390)
text('NÜMUNƏ / DATA MƏRKƏZİ',806,216,300,12,GREEN,True)
text('4 server\n2 kommutator',806,252,299,29,INK,True)
line(806,347,1106,347)
text('57 000 AZN',806,374,299,34,GREEN,True)
text('İlkin smeta, ƏDV-siz. Məbləğlər mövqelər əsasında serverdə hesablanır.',806,434,292,19,MUTED)
end()

# 07
start('Daxili nəzarət','Təsdiq prosesi aydın qaydalara əsaslanır',
      'Məbləğə görə təsdiq zənciri və vəzifələrin ayrılması.',
      'Müəllif öz tələbini təsdiqləyə bilməz; bir şəxs eyni tələbin iki təsdiq addımını bağlaya bilməz. Şəkildəki hədlər demo konfiqurasiyasıdır, universal hüquqi limit deyil.',
      [S+'requests/requests.module.ts:177–201','scripts/produce-guides.mjs:462–509'])
shot('buyer-settings-94s.png',46,208,1060,300,crop=(264,320,1248,680))
text('Müəllif ≠ təsdiqləyən',46,550,330,21,GREEN,True)
text('Məbləğə görə marşrut',410,550,335,21,GREEN,True)
text('Hər addımın tarixçəsi',774,550,330,21,GREEN,True)
end()

# 08
start('Satınalma üsulları','Hər ehtiyac üçün uyğun prosedur',
      'Metod siyasəti kateqoriya və məbləğ hədlərinə əsaslanır.',
      'RFI məlumat toplamaq üçündür; qiymət üzrə qalib seçimi və müqavilə mükafatı üçün nəzərdə tutulmur. Birbaşa sifarişin uyğunluğu təşkilat siyasətindən asılıdır.',
      ['apps/web/src/guides.ts:66–74',S+'sourcing/sourcing.service.ts:349–355'])
methods=[('RFQ','Qiymət sorğusu','Müəyyən edilmiş tələblər üzrə məxfi qiymət təklifləri.'),('RFP','Texniki təklif','Qiymətlə yanaşı məzmunlu texniki təklifin qiymətləndirilməsi.'),('RFI','Məlumat sorğusu','Bazar və həll variantları haqqında qiymətsiz məlumat toplama.'),('AUKSİON','Azalan qiymət','Təchizatçıların qaydalı şəkildə daha aşağı qiymət verməsi.'),('BİRBAŞA','Sadə satınalma','Siyasətin icazə verdiyi məbləğlər üçün birbaşa PO.'),('KATALOQ','Təkrar sifariş','Razılaşdırılmış qiymətlər və minimum sifariş miqdarı.')]
for i,(n,t,b) in enumerate(methods):
    x=46+i%3*362; y=210+i//3*185
    rect(x,y,336,163,WHITE,LINE,r=7); text(n,x+20,y+15,295,12,GREEN,True); text(t,x+20,y+43,295,22,INK,True); text(b,x+20,y+82,295,17,MUTED,max_h=76)
end()

# 09
start('Məxfi tender','Qiymətlər açılışa qədər qorunur',
      'Alıcı iştirak faktını görür; təchizatçı öz təklifini idarə edir.',
      'Bu ekran TƏCHİZATÇI kabinetidir: öz qiymətini görməsi normaldır. Alıcı məbləği açılışa qədər görə bilmir. Təklifin məbləği şifrələnir; ümumi sual-cavab bütün uyğun iştirakçılar üçün eyni məlumatı təmin edir.',
      ['apps/api/src/common/crypto/sealed.service.ts:5–33',S+'sourcing/sourcing.service.ts:327–415','scripts/produce-guides.mjs:520–545'])
shot('sup-market-43s.png',46,211,737,380)
text('TƏCHİZATÇI KABİNETİ',818,218,288,12,GREEN,True)
feature('01','Şifrələnmiş məbləğ','Açılışdan əvvəl alıcı üçün qiymət məlumatı bağlıdır.',818,256,288)
feature('02','Bərabər məlumat','Sual-cavab bölməsində cavablar iştirakçılar üçün ortaqdır.',818,426,288)
end()

# 10
start('Nümayiş nümunəsi','57 000 AZN smetadan 51 500 AZN təklifə',
      'Data mərkəzi üçün 4 server və 2 kommutatorun satınalması.',
      'Rəqəmlər scripts/demo-fill.mjs-dəki nümayiş ssenarisindəndir. Bu, real müştəri nəticəsi və ya qənaət zəmanəti deyil. Qiymətlər ƏDV-sizdir. 5 500 / 57 000 = 9,6491%, iki onluğa yuvarlaq 9,65%.',
      ['scripts/demo-fill.mjs:197–246'])
cols=[46,374,617,861]; widths=[328,243,244,245]
heads=['Mövqe','İlkin smeta','Supply One','Baku Metal']
for x,w,t in zip(cols,widths,heads): rect(x,210,w,48,FOREST); text(t,x+16,222,w-25,16,WHITE,True)
rows=[['4 server','48 000 AZN','43 200 AZN','46 000 AZN'],['2 kommutator','9 000 AZN','8 300 AZN','8 600 AZN'],['Cəmi / ƏDV-siz','57 000 AZN','51 500 AZN','54 600 AZN'],['Çatdırılma','—','21 gün','14 gün']]
for i,row in enumerate(rows):
    for j,(x,w,t) in enumerate(zip(cols,widths,row)):
        rect(x,258+i*47,w,47,MINT if i==2 else WHITE,LINE)
        text(t,x+16,271+i*47,w-26,18,GREEN if i==2 else INK,i==2)
rect(46,477,1060,90,MINT,r=8)
text('5 500 AZN',69,493,250,35,GREEN,True); text('9,65%',352,493,180,35,GREEN,True)
text('Smeta ilə qalib təklif arasındakı fərq.\nSeçim qiymət və çatdırılma meyarları ilə əsaslandırılır.',567,495,515,19,FOREST)
text('Nümayiş məlumatlarıdır; real müştəri nəticəsi deyil.',47,582,1000,11,MUTED)
end()

# 11
start('Qiymətləndirmə və qərar','Qalib seçiminin məntiqi görünür',
      'Əvvəlcədən müəyyən edilən meyarlar, çəkili bal və lot üzrə qərar.',
      'Demoda meyarlar qiymət 70%, çatdırılma 30%-dir; Supply One 95×0,7+75×0,3=89; Baku Metal 82×0,7+90×0,3=84,4. AI tövsiyəsi ayrıca dəstək funksiyasıdır, bu cədvəldəki rəsmi qiymətləndirmənin əvəzi deyil. Yekun qərar alıcıdadır.',
      [S+'sourcing/sourcing.service.ts:485–618','scripts/demo-fill.mjs:218–230','services/ai-orchestrator/main.py:68–114'])
shot('buyer-events-28s.png',46,205,1060,277,crop=(265,245,1247,539))
feature('01','Meyarlar əvvəlcədən','Tender yaradılarkən meyarlar sabitlənir.',46,514,330)
feature('02','İstisna əsaslandırılır','Reytinqdən fərqli seçim üçün səbəb qeydə alınır.',408,514,330)
feature('03','İzah edilən tövsiyə','Qiymət, müddət və tarixçə qərarı dəstəkləyir.',770,514,336)
end()

# 12
start('Əsas imkan / revers-auksion','Rəqabət qiyməti addım-addım aşağı salır',
      'Ən yaxşı təklif, növbəti qiymət həddi və son tarix bir paneldə.',
      'Auksion nümunəsi dizel yanacağı üçündür. Şəkildə start 32 000, ən yaxşı qiymət 31 200, növbəti maksimum 30 880 AZN-dir. Son saniyədə təklif veriləndə qaydaya əsasən vaxt uzadılır. Nəticə nümunədir.',
      [S+'sourcing/auction.service.ts:98–111,164–200','scripts/produce-guides.mjs:573–584'])
shot('sup-auction-14s.png',46,205,1060,260,crop=(263,250,1247,480))
for i,(n,t,b) in enumerate([('01','Minimum addım','Yeni təklif cari ən yaxşı qiymətdən müəyyən addım aşağı olmalıdır.'),('02','Vaxtın uzadılması','Son saniyə təklifləri üçün avtomatik uzatma qaydası.'),('03','Rəqabətin görünməsi','Təchizatçı cari ən yaxşı qiyməti izləyir.')]): feature(n,t,b,46+i*362,499,336)
end()

# 13
start('Təchizatçı iştirakı','Dəvətdən cavaba qədər izləmə avtomatlaşdırılır',
      'Təklif verməmiş təchizatçılara planlı xatırlatmalar.',
      'SMTP və WhatsApp/Meta kanalları ayrıca sazlanır. Təklif verildikdə xatırlatma zənciri dayanır. Burada canlı mesaj göndərilməsi yoxlanılmayıb; göstərilən demo planıdır.',
      [S+'followups/followups.module.ts:186–249',S+'channels/channels.module.ts:30–79',S+'sourcing/sourcing.service.ts:413–415'])
shot('buyer-events-84s.png',46,211,1060,290,crop=(264,315,1249,614))
text('Dəvət',46,550,160,23,FOREST,True); text('→',210,550,40,23,GREEN)
text('Planlı xatırlatma',271,550,260,23,FOREST,True); text('→',558,550,40,23,GREEN)
text('Təklif daxil olur',630,550,250,23,FOREST,True); text('→',904,550,40,23,GREEN)
text('Zəncir dayanır',960,550,145,20,FOREST,True)
end()

# 14
start('Müqavilə və icra','Qalib təklif icra edilən sifarişə çevrilir',
      'Müqavilə, qiymət şərtləri, PO və çatdırılma mərhələləri əlaqəlidir.',
      'Müqavilə qalib tender əsasında yaranır. PO-nun təsdiq, qablaşdırma, göndəriş, çatdırılma və qəbul mərhələləri var. Logistika hadisələrini təchizatçı daxil edir; GPS və ya daşıyıcı inteqrasiyası iddia edilmir.',
      ['apps/web/src/guides.ts:110–129',S+'logistics/logistics.module.ts'])
shot('buyer-contracts.az.png',46,206,1060,309,crop=(262,0,1252,480))
text('Müqavilə',46,553,220,23,GREEN,True); text('→',280,553,40,23,GREEN)
text('PO',359,553,150,23,GREEN,True); text('→',541,553,40,23,GREEN)
text('Çatdırılma',618,553,220,23,GREEN,True); text('→',860,553,40,23,GREEN)
text('Qəbul',960,553,146,23,GREEN,True)
end()

# 15
start('Dəyişiklik nəzarəti','Sifarişdə dəyişiklik qarşı tərəfin razılığı ilə',
      'Köhnə və yeni şərtlər görünür; tarixçə saxlanılır.',
      'İkinci nümayiş nümunəsi: A4 kağız 40×7=280 AZN-dən 60×7=420 AZN-ə dəyişir. Təchizatçı qəbul və ya rədd edir. Məbləğ yalnız qəbuldan sonra yenilənir. Sifariş dəyişiklikləri üçün status və qəbul şərtləri mövcuddur.',
      ['scripts/demo-fill.mjs:307–315','scripts/produce-guides.mjs:547–558'])
shot('sup-orders-18s.png',46,210,1060,280,crop=(264,144,1245,440))
rect(46,527,1060,60,MINT,r=6)
text('40 → 60 bağlama',68,542,320,25,FOREST,True)
text('280 → 420 AZN',420,542,320,25,FOREST,True)
text('Qəbul / rədd',838,542,240,25,FOREST,True)
end()

# 16
start('Maliyyə nəzarəti','Sifariş, qəbul və qaimə üçtərəfli yoxlanılır',
      'Fərqlər görünür; təsdiq qərarı faktiki əməliyyata söykənir.',
      'Sistem sifariş, qəbul edilmiş miqdar və qaimə sətirlərini üzləşdirir. Uyğunsuzluq DISPUTED statusu yaradır. Qəbul sonradan daxil edilərsə, yoxlama yenilənir. PAID yalnız ödənişin uçotudur, bank əməliyyatı deyil. Demo ƏDV hesabı 51 500×18%=9 270; cəmi 60 770.',
      [S+'invoices/invoices.module.ts:32–84,214–220,292–295',S+'contracts/contracts.module.ts:320–327','scripts/demo-fill.mjs:237–246'])
for i,(t,b) in enumerate([('Sifariş (PO)','Nə və hansı qiymətə alınıb?'),('Qəbul','Faktiki nə qədər qəbul edilib?'),('Qaimə','Qaimədə nə göstərilib?')]):
    x=46+i*362; rect(x,208,336,100,MINT,r=7); text(t,x+20,224,298,24,FOREST,True); text(b,x+20,267,298,16,MUTED)
shot('buyer-invoices.az.png',46,337,1060,140,crop=(263,0,1247,210))
text('51 500 + 9 270 = 60 770 AZN',47,524,705,31,GREEN,True)
text('Məbləğ + 18% ƏDV = qaimənin yekunu · demo',47,569,760,15,MUTED)
text('Ödəniş statusu\nsistemdə qeydə alınır.',840,530,265,20,INK,True)
end()

# 17
start('Təkrar satınalma','Razılaşdırılmış qiymətlərlə təkrar sifariş verin',
      'Kataloqlar standart mallar üzrə təkrar sifarişi sadələşdirir.',
      'Aktiv kataloq müqavilədən yaradıla və ya razılaşdırılmış qiymətləri saxlaya bilər. Sifarişdə qiymət və minimum miqdar yoxlanılır. Bunlar təşkilatın münasibət və lisenziya qaydalarına tabedir.',
      ['apps/web/src/guides.ts:143–153',S+'catalogs/catalogs.module.ts'])
shot('buyer-catalogs.az.png',46,218,1060,181,crop=(263,0,1247,220))
for i,(n,t,b) in enumerate([('01','Razılaşdırılmış qiymət','Eyni malı hər dəfə yenidən qiymətləndirmək ehtiyacı azalır.'),('02','Minimum miqdar','Kataloqdakı sifariş hədləri avtomatik yoxlanılır.'),('03','Birbaşa PO','Uyğun kataloq mövqeləri üzrə satınalma sifarişi yaradılır.')]): feature(n,t,b,46+i*362,450,336)
end()

# 18
start('Büdcə və analitika','Plan, öhdəlik və faktiki xərc ayrı görünür',
      'Maliyyə nəzarəti və idarəetmə hesabatı eyni məlumat zəncirinə əsaslanır.',
      'Büdcə istifadəsi təsdiqlənmiş tələblərin öhdəliyi əsasında hesablanır, analitikada xərc isə ACCEPTED PO-lardır. Aşım xəbərdarlığı avtomatik bloklama deyil. Demo ekranlarındakı qənaət real müştəri nəticəsi sayılmır.',
      [S+'budgets/budgets.module.ts:30–63,114–137',S+'dashboard/dashboard.module.ts:86–147'])
text('Büdcə və öhdəlik',46,207,508,23,INK,True)
text('Faktiki xərc və qənaət',598,207,508,23,INK,True)
shot('buyer-budgets.az.png',46,250,508,235,crop=(264,24,1248,470))
shot('buyer-analytics.az.png',598,250,508,292,crop=(263,30,1248,644))
text('Limit aşımı barədə xəbərdarlıq.\nYekun qərar səlahiyyətli şəxsdədir.',46,526,508,19,MUTED)
end()

# 19
start('Təchizatçı idarəçiliyi','Təchizatçı bazasını sənəd və nəticə ilə qiymətləndirin',
      'Reyestr, sənəd yoxlaması və icra göstəriciləri bir-birini tamamlayır.',
      'KYC burada sənədlərin toplanması və alıcının qərarıdır; dövlət reyestrinə avtomatik sorğu deyil. Reytinq vaxtında çatdırılma, keyfiyyət və kommunikasiya qiymətlərini toplayır.',
      ['apps/web/src/guides.ts:207–232',S+'verification/verification.module.ts'])
shot('buyer-suppliers.az.png',46,210,1060,285,crop=(263,0,1249,409))
for i,(n,t,b) in enumerate([('01','Reyestr','Dəvət, təsdiq və bloklama.'),('02','Sənəd yoxlaması','VÖEN, lisenziya və tələb edilən sənədlər.'),('03','Təchizatçı reytinqi','Çatdırılma, keyfiyyət və ünsiyyət.')]): feature(n,t,b,46+i*362,522,336)
end()

# 20
start('Etibar və uyğunlaşdırma','Korporativ tələblərə uyğun iş mühiti',
      'Məlumat nəzarəti, lokal iş qaydaları və müxtəlif yerləşdirmə variantları.',
      'RLS/ACL, AES-256-GCM və append-only audit kod səviyyəsində nəzərdə tutulmuş nəzarətlərdir. Sertifikasiya, penetrasiya testinin yeni nəticəsi və ya canlı SLA iddia edilmir. Yerləşdirmə və funksiyalar lisenziya və texniki konfiqurasiyadan asılıdır.',
      ['README.md:3–33','infra/DEPLOY.md:1–63','apps/api/src/common/crypto/sealed.service.ts:5–33','apps/api/src/common/audit/audit.service.ts:6–43'])
card('MƏLUMAT','Ayrılmış giriş','Təşkilatlar üzrə məlumat sərhədləri, rol əsaslı səlahiyyət və audit tarixçəsi.',46,210,340,246)
card('LOKAL İŞ MÜHİTİ','AZ · RU · EN','Azərbaycan dili, AZN, ƏDV və VÖEN; bölmələr üzrə video bələdçilər.',406,210,340,246)
card('YERLƏŞDİRMƏ','Bulud və on-premise','Ortaq bulud modeli və müəssisənin öz infrastrukturunda yerləşdirmə seçimi.',766,210,340,246)
rect(46,500,1060,80,MINT,r=7)
text('Audit üçün izlənən qərarlar. Komanda üçün aydın rollar. Təchizatçı üçün vahid portal.',68,521,1016,23,FOREST,True)
end()

# 21
start('Məlumat mübadiləsi','Məlumat mübadiləsi və qoşulma imkanları',
      'Hazır ixrac imkanları və ayrıca qurulan bağlantılar.',
      'OCDS ixracı tender və qalib məlumatları üçündür. UBL faylı Peppol şəbəkəsi vasitəsilə öz-özünə göndərilmir; Access Point tələb olunur. 1C üçün CSV var, iki tərəfli sinxronizasiya deyil. ASAN İmza operator inteqrasiyası tələb edir; hazırda development provider işləyir.',
      [S+'sourcing/sourcing.service.ts:674–712',S+'invoices/invoices.module.ts:119–148',S+'channels/channels.module.ts:30–79','apps/api/src/common/signature/signature.provider.ts:5–35'])
rows=[('OCDS','Tender və qalib məlumatlarının strukturlaşdırılmış ixracı.'),('UBL / qaimə','Elektron qaimə faylı; Peppol ötürülməsi üçün ayrıca operator lazımdır.'),('1C / CSV','Qaimələrin 1C-yə uyğunlaşdırıla bilən CSV faylı ilə ixracı.'),('E-poçt / WhatsApp','Müştərinin kanalları və təsdiqlənmiş şablonları ilə konfiqurasiya olunur.'),('ASAN İmza','Hazırda aktiv deyil; operator bağlantısı ayrıca qurulmalıdır.')]
for i,(a,b) in enumerate(rows):
    y=206+i*73; rect(46,y,1060,61,WHITE,LINE,r=5); text(a,66,y+18,235,20,GREEN,True); text(b,315,y+17,770,18,INK,max_h=50)
end()

# 22
start('Növbəti addım','Bir pilot satınalma ilə dəyəri ölçün',
      'Real ehtiyac, seçilmiş təchizatçılar və əvvəlcədən razılaşdırılmış uğur meyarları.',
      'Bu, təklif olunan pilot planıdır; müqavilə öhdəliyi, qiymət və ya müddət zəmanəti deyil. Qiymətləndirmə meyarları müştəri ilə əvvəlcədən razılaşdırılmalıdır.',
      ['Təklif olunan müştəri pilot planı — müəllif tövsiyəsi'])
for i,(n,t,b) in enumerate([('01','Çərçivəni seçin','Bir kateqoriya, bir satınalma və məsul komanda.'),('02','Qaydaları qurun','Rollar, təsdiq hədləri, qiymətləndirmə və təchizatçılar.'),('03','Dövrü tamamlayın','Tələbdən təklifə, sifarişdən qəbul və qaiməyə.'),('04','Nəticəni qiymətləndirin','Dövr müddəti, iştirak, smeta fərqi və uyğunsuzluqlar.')]):
    x=46+i*270; rect(x,222,250,230,WHITE,LINE,r=8); text(n,x+20,242,210,15,GREEN,True); text(t,x+20,286,210,25,INK,True,max_h=72); text(b,x+20,365,210,17,MUTED,max_h=80)
rect(46,496,1060,87,FOREST,r=8)
text('Növbəti görüş: şirkətinizin bir satınalma ssenarisini birlikdə xəritələndirək.',70,518,1010,27,WHITE,True)
end()

prs.save(OUT/'Leadprocure_Korporativ_Teqdimat_AZ.pptx'); pdf.save()
(ROOT/'slide-content.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
(OUT/'qa-layout.json').write_text(json.dumps({'slides':len(manifest),'issues':violations,'source_count':len(all_sources)},ensure_ascii=False,indent=2))
notes=['# Leadprocure — məruzəçi qeydləri\n']
for m in manifest: notes += [f"## {m['slide']:02}. {m['title']}\n",m['notes']+'\n','Mənbələr: '+ '; '.join(m['sources'])+'\n']
(OUT/'Meruzeci_Qeydleri_AZ.md').write_text('\n'.join(notes))
doc=fitz.open(OUT/'Leadprocure_Korporativ_Teqdimat_AZ.pdf')
renders=OUT/'preview'; renders.mkdir(exist_ok=True)
for i,p in enumerate(doc):
    p.get_pixmap(matrix=fitz.Matrix(1,1),alpha=False).save(renders/f'slide-{i+1:02}.png')
from PIL import ImageDraw
for start_i in range(0,len(doc),6):
    sheet=Image.new('RGB',(1200,1092),'#e4e6e0'); d=ImageDraw.Draw(sheet)
    for j,i in enumerate(range(start_i,min(start_i+6,len(doc)))):
        im=Image.open(renders/f'slide-{i+1:02}.png'); im.thumbnail((576,324)); x=16+j%2*600; y=22+j//2*364; sheet.paste(im,(x,y)); d.text((x,y+327),f"{i+1:02} / {manifest[i]['title']}",fill='#202b27')
    sheet.save(OUT/f'contact-sheet-{start_i//6+1}.jpg',quality=90)
print(json.dumps({'slides':len(manifest),'layout_issues':violations,'pptx_bytes':(OUT/'Leadprocure_Korporativ_Teqdimat_AZ.pptx').stat().st_size,'pdf_bytes':(OUT/'Leadprocure_Korporativ_Teqdimat_AZ.pdf').stat().st_size},ensure_ascii=False))
